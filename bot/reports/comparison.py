"""
Генерация сравнительного Excel-отчёта по двум магазинам.
"""

import re
import logging

from openpyxl import Workbook
from openpyxl.styles import Font

from bot.config import THRESHOLD_A, THRESHOLD_B, THRESHOLD_C
from bot.services.calculations import aggregate_by_article
from bot.reports.excel_styles import (
    DATA_FONT_BOLD, CENTER,
    CMP_RAISE_BG, CMP_LOWER_BG, CMP_PAIR_ALT_BG,
    fill,
    apply_header_style, write_legend_block,
    write_merged_article, apply_row_background, write_store_row,
    make_report_path,
)

logger = logging.getLogger(__name__)


def generate_comparison_report(
    store1_data: list[dict],
    store2_data: list[dict],
    store1_name: str,
    store2_name: str,
    days_threshold: int = 7,
    threshold_a: float = THRESHOLD_A,
    threshold_b: float = THRESHOLD_B,
    threshold_c: float = THRESHOLD_C,
) -> str | None:
    """
    Генерирует сравнительный Excel-отчёт по двум магазинам.

    Сопоставляет товары по supplier_article. Для каждой пары показывает
    данные обоих магазинов и рекомендацию по цене.

    Returns:
        Путь к файлу или None если нет совпадающих позиций.
    """
    logger.info(f"=== Сравнительный отчёт: {store1_name} vs {store2_name} ===")

    agg1 = aggregate_by_article(store1_data)
    agg2 = aggregate_by_article(store2_data)

    # Общие артикулы
    common_articles = sorted(set(agg1.keys()) & set(agg2.keys()))
    only_s1 = len(set(agg1.keys()) - set(agg2.keys()))
    only_s2 = len(set(agg2.keys()) - set(agg1.keys()))

    logger.info(f"Совпадений: {len(common_articles)}, только в маг.1: {only_s1}, только в маг.2: {only_s2}")

    if not common_articles:
        return None

    # Создаём Excel вручную (для merged cells)
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение"

    # Шапка
    headers = ['Артикул', 'Магазин', 'ID (WB)', 'Остаток', 'Дней осталось',
               'Продаж/день', 'Группа', 'Цена ₽', 'Рекомендация']
    col_widths = {1: 22, 2: 24, 3: 15, 4: 12, 5: 16, 6: 15, 7: 11, 8: 13, 9: 18}

    apply_header_style(ws, col_widths, headers=headers, auto_filter=False)

    num_cols = len(headers)
    current_row = 2

    for pair_idx, article in enumerate(common_articles):
        s1 = agg1[article]
        s2 = agg2[article]

        # Определяем рекомендацию
        avg1 = s1.get('avg_per_day') or 0
        avg2 = s2.get('avg_per_day') or 0
        max_avg = max(avg1, avg2)

        if max_avg == 0:
            rec1, rec2 = "—", "—"
        elif abs(avg1 - avg2) / max_avg < 0.1:
            rec1, rec2 = "—", "—"
        elif avg1 > avg2:
            rec1, rec2 = "↑ Повысить", "↓ Понизить"
        else:
            rec1, rec2 = "↓ Понизить", "↑ Повысить"

        row1 = current_row
        row2 = current_row + 1

        # Чередование фона пар
        pair_bg = CMP_PAIR_ALT_BG if pair_idx % 2 == 1 else None

        for row_num, store_data, store_name_val, rec, is_top in [
            (row1, s1, store1_name, rec1, True),
            (row2, s2, store2_name, rec2, False),
        ]:
            is_bottom = not is_top

            # B-H: общие колонки
            border_fn = write_store_row(
                ws, row_num, store_data, store_name_val,
                is_top, is_bottom, num_cols,
            )

            # I: Рекомендация
            c = ws.cell(row=row_num, column=9, value=rec)
            c.font = DATA_FONT_BOLD
            c.alignment = CENTER
            c.border = border_fn(9)
            if "Повысить" in rec:
                c.fill = fill(CMP_RAISE_BG)
            elif "Понизить" in rec:
                c.fill = fill(CMP_LOWER_BG)

            # Фон чередования
            if pair_bg:
                apply_row_background(ws, row_num, pair_bg, num_cols)

        # A: Артикул (merged)
        write_merged_article(ws, article, row1, row2, num_cols)

        current_row += 2

    # Итоги внизу
    summary_row = current_row + 1
    summary_font = Font(name="Arial", size=9, bold=True, color="555555")

    ws.cell(row=summary_row, column=1, value=f"Совпадающих позиций: {len(common_articles)}").font = summary_font
    ws.cell(row=summary_row + 1, column=1, value=f"Только в {store1_name}: {only_s1}").font = summary_font
    ws.cell(row=summary_row + 2, column=1, value=f"Только в {store2_name}: {only_s2}").font = summary_font

    # Легенда
    legend_row = summary_row + 4
    write_legend_block(ws, legend_row, [
        ("— Группы товаров (продаж/день) —", "title"),
        (f"A: ходовые (≥{threshold_a})", "item"),
        (f"B: средние (≥{threshold_b})", "item"),
        (f"C: редкие (≥{threshold_c})", "item"),
        (f"D: почти не продаются (<{threshold_c})", "item"),
        ("", "item"),
        ("— Рекомендации —", "title"),
        ("↑ Повысить — товар продаётся лучше конкурента (можно повысить цену)", "item"),
        ("↓ Понизить — товар отстаёт (снизить цену для стимуляции спроса)", "item"),
        ("— паритет (разница продаж < 10%)", "item"),
    ])

    # Сохранение
    def _short_name(full_name):
        if not full_name:
            return "store"
        short = full_name.split('(')[0].strip()
        if not short:
            short = full_name.split()[0] if full_name.split() else "store"
        return re.sub(r'[^\w-]', '', short).strip()[:20]

    output_path = make_report_path(
        f'comparison_{_short_name(store1_name)}_{_short_name(store2_name)}'
    )

    wb.save(output_path)
    logger.info(f"✓ Сравнительный отчёт: {output_path}")
    return output_path
