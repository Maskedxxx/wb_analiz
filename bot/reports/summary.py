"""
Генерация сводного Excel-отчёта по всем магазинам.
"""

import logging

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from bot.config import THRESHOLD_A, THRESHOLD_B, THRESHOLD_C
from bot.services.calculations import aggregate_by_article, merge_wh_by_name
from bot.reports.excel_styles import (
    DATA_FONT, DATA_FONT_BOLD, CENTER, LEFT, WRAP_LEFT,
    HYPERLINK_FONT, SUMMARY_PAIR_ALT_BG,
    GROUP_COLORS, WB_CABINET_URL,
    fill,
    apply_header_style, write_legend_block,
    write_merged_article, apply_row_background, write_store_row,
    add_stock_comment, group_border_fn, make_report_path,
)

logger = logging.getLogger(__name__)


def generate_summary_report(
    all_stores_data: dict,
    all_warehouse_data: dict,
    days_threshold: int = 7,
    threshold_a: float = THRESHOLD_A,
    threshold_b: float = THRESHOLD_B,
    threshold_c: float = THRESHOLD_C,
) -> str | None:
    """
    Генерирует сводный Excel-отчёт по всем магазинам.

    Показывает только товары (по supplier_article), присутствующие в 2+ магазинах.
    Вместо рекомендации — детализация остатков по складам.

    Args:
        all_stores_data: {store_name: [product_rows]}
        all_warehouse_data: {store_name: [{nm_id, warehouse_name, quantity}]}

    Returns:
        Путь к файлу или None если нет совпадающих позиций.
    """
    logger.info(f"=== Сводный отчёт: {len(all_stores_data)} магазинов ===")

    # Агрегация по артикулу для каждого магазина
    aggregated = {}
    for store_name, rows in all_stores_data.items():
        aggregated[store_name] = aggregate_by_article(rows)

    # Маппинг nm_id → supplier_article из product data
    nm_to_article = {}
    for store_name, rows in all_stores_data.items():
        for r in rows:
            nm_to_article.setdefault(store_name, {})[r['nm_id']] = r.get('supplier_article')

    # Дополняем маппинг из warehouse data (для nm_id без заказов)
    for store_name, wh_rows in all_warehouse_data.items():
        store_nm_map = nm_to_article.setdefault(store_name, {})
        for wr in wh_rows:
            if wr['nm_id'] not in store_nm_map:
                sa = wr.get('supplier_article')
                if sa:
                    store_nm_map[wr['nm_id']] = sa

    # Построение индекса складов по article (не nm_id)
    wh_index = {}  # {store_name: {article: [{warehouse_name, quantity, in_way_from_client}]}}
    for store_name, wh_rows in all_warehouse_data.items():
        store_wh = {}
        store_nm_map = nm_to_article.get(store_name, {})
        for wr in wh_rows:
            article = store_nm_map.get(wr['nm_id'])
            if not article:
                continue
            store_wh.setdefault(article, []).append({
                'warehouse_name': wr['warehouse_name'],
                'quantity': wr['quantity'],
                'in_way_from_client': wr.get('in_way_from_client', 0),
            })
        wh_index[store_name] = store_wh

    # Найти артикулы, присутствующие в 2+ магазинах
    article_stores = {}  # {article: [store_name, ...]}
    for store_name, agg in aggregated.items():
        for article in agg:
            article_stores.setdefault(article, []).append(store_name)

    common_articles = sorted([
        art for art, stores in article_stores.items() if len(stores) >= 2
    ])

    logger.info(f"Артикулов в 2+ магазинах: {len(common_articles)}")

    if not common_articles:
        return None

    # Создаём Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводный отчёт"

    # Шапка
    headers = ['Артикул', 'Магазин', 'ID (WB)', 'Остаток', 'Дней осталось',
               'Продаж/день', 'Группа', 'Цена ₽', 'Остатки по складам']
    col_widths = {1: 22, 2: 24, 3: 15, 4: 12, 5: 16, 6: 15, 7: 11, 8: 13, 9: 36}

    apply_header_style(ws, col_widths, headers=headers, auto_filter=False)

    num_cols = len(headers)
    current_row = 2

    for group_idx, article in enumerate(common_articles):
        stores_with_article = article_stores[article]
        group_size = len(stores_with_article)
        group_start = current_row

        # Чередование фона
        group_bg = SUMMARY_PAIR_ALT_BG if group_idx % 2 == 1 else None

        for store_idx, store_name in enumerate(stores_with_article):
            row_num = current_row
            is_top = (store_idx == 0)
            is_bottom = (store_idx == group_size - 1)
            store_data = aggregated[store_name][article]

            # B-H: общие колонки (без фиксации высоты — Excel подберёт под wrap_text)
            border_fn = write_store_row(
                ws, row_num, store_data, store_name,
                is_top, is_bottom, num_cols,
                row_height=None,
            )

            # Warehouse data для колонок D (перезапись) и I
            wh_list = wh_index.get(store_name, {}).get(article, [])
            wh_list = merge_wh_by_name(wh_list)
            total_iwfc = sum(w.get('in_way_from_client', 0) for w in wh_list)
            wh_active = [w for w in wh_list if w['quantity'] > 0]
            wh_active.sort(key=lambda w: w['quantity'], reverse=True)

            # D: Перезапись остатка в формат "qty (+возвр)"
            if wh_active:
                wh_total_qty = sum(w['quantity'] for w in wh_active)
                display_stock = f"{wh_total_qty} (+{total_iwfc})" if total_iwfc > 0 else str(wh_total_qty)
            else:
                stock_val = store_data.get('stock_qty', 0)
                iwfc_total = store_data.get('in_way_from_client', 0)
                display_stock = f"{stock_val} (+{iwfc_total})" if iwfc_total > 0 else str(stock_val)
            cell_d = ws.cell(row=row_num, column=4, value=display_stock)
            cell_d.font = DATA_FONT
            cell_d.alignment = CENTER
            cell_d.border = border_fn(4)

            # Комментарий к остатку
            add_stock_comment(cell_d, wh_list)

            # I: Остатки по складам — компактная строка
            if wh_active:
                compact_parts = [f"{w['warehouse_name']}: {w['quantity']}" for w in wh_active]
                if total_iwfc > 0:
                    compact_parts.append(f"+{total_iwfc} возвр.")
                compact = " | ".join(compact_parts)
            else:
                stock_val = store_data.get('stock_qty', 0)
                iwfc_total = store_data.get('in_way_from_client', 0)
                parts = []
                if stock_val > 0:
                    parts.append(f"итого: {stock_val}")
                if iwfc_total > 0:
                    parts.append(f"+{iwfc_total} возвр.")
                compact = " | ".join(parts) if parts else "—"

            c = ws.cell(row=row_num, column=9, value=compact)
            c.font = Font(name="Arial", size=9, color="555555")
            c.alignment = WRAP_LEFT
            c.border = border_fn(9)

            # Фон чередования
            if group_bg:
                apply_row_background(ws, row_num, group_bg, num_cols)

            current_row += 1

        # A: Артикул (merged)
        write_merged_article(ws, article, group_start, group_start + group_size - 1, num_cols)

    # Итоги
    summary_row = current_row + 1
    summary_font = Font(name="Arial", size=9, bold=True, color="555555")
    total_articles = len(common_articles)
    total_stores = len(all_stores_data)

    ws.cell(row=summary_row, column=1,
            value=f"Общих позиций (в 2+ магазинах): {total_articles}").font = summary_font
    ws.cell(row=summary_row + 1, column=1,
            value=f"Магазинов в отчёте: {total_stores}").font = summary_font

    # Легенда
    legend_row = summary_row + 3
    write_legend_block(ws, legend_row, [
        ("— Группы товаров (продаж/день) —", "title"),
        (f"A: ходовые (≥{threshold_a})", "item"),
        (f"B: средние (≥{threshold_b})", "item"),
        (f"C: редкие (≥{threshold_c})", "item"),
        (f"D: почти не продаются (<{threshold_c})", "item"),
        ("", "item"),
        ("— Остатки по складам —", "title"),
        ("Наведите мышь на ячейку для детализации по складам", "item"),
    ])

    # Сохранение
    output_path = make_report_path('summary')

    wb.save(output_path)
    logger.info(f"✓ Сводный отчёт: {output_path}")
    return output_path
