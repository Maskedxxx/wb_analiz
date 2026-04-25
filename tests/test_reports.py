"""
Тесты генерации отчётов: проверка что Excel-файлы создаются корректно.
"""

import os

import pytest
from openpyxl import load_workbook

from bot.reports.single import generate_report_from_data
from bot.reports.comparison import generate_comparison_report
from bot.reports.summary import generate_summary_report


class TestSingleReport:
    def test_generates_file(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test Store', warehouse_rows=warehouse_rows,
        )
        assert path is not None
        assert os.path.exists(path)

    def test_sheets(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
        )
        wb = load_workbook(path)
        assert set(wb.sheetnames) == {'Все товары', 'Поставки - расчёт', 'Поставки - предложение'}

    def test_all_products_sheet_has_all_rows(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
        )
        wb = load_workbook(path)
        ws = wb['Все товары']
        data_rows = [r for r in ws.iter_rows(min_row=2) if r[0].value is not None]
        assert len(data_rows) == len(product_rows)

    def test_empty_data(self, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data([], store_name='Empty')
        assert os.path.exists(path)

    def test_refill_calc_sheet_uses_if_formula(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        """Лист «Поставки - расчёт» с warehouse_distribution: колонки D..K содержат IF-формулы."""
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        dist = [
            {'warehouse_id': 507, 'weight': 50, 'cutoff_days': 3},
        ]
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
            warehouse_distribution=dist,
        )
        wb = load_workbook(path)
        ws = wb['Поставки - расчёт']
        # Row 4 — шапка, данные с row 5. ART-1 имеет avg>0 → попадает в данные.
        d_cell = ws.cell(row=5, column=4)
        assert isinstance(d_cell.value, str) and d_cell.value.startswith('=IF(')

    def test_refill_calc_hidden_columns(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        """Лист «Поставки - расчёт»: колонки M..U скрыты (days_cover + stock per wh)."""
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
        )
        wb = load_workbook(path)
        ws = wb['Поставки - расчёт']
        assert ws.column_dimensions['M'].hidden is True

    def test_wb_suggestion_headers(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        """Шапка листа «Поставки - предложение» содержит ожидаемые заголовки."""
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
        )
        wb = load_workbook(path)
        ws = wb['Поставки - предложение']
        assert ws.cell(row=1, column=1).value == 'Артикул'
        assert ws.cell(row=1, column=2).value == 'Баркод'
        assert ws.cell(row=1, column=3).value == 'Объём WB'
        assert ws.cell(row=1, column=4).value == 'Оборотность'
        assert ws.cell(row=1, column=6).value == 'Упущено заказов'
        assert ws.cell(row=1, column=7).value == 'Срок продаж (WB)'
        assert ws.cell(row=1, column=8).value == 'Тренд'

    def test_nonliquid_zero_volume(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        """Товар с availability=nonLiquid → объём WB = 0 на листе «Поставки - предложение»."""
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
        )
        wb = load_workbook(path)
        ws = wb['Поставки - предложение']
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == 'ART-2':
                assert ws.cell(row=row, column=3).value == 0
                assert ws.cell(row=row, column=4).value == 'неликвид'
                return
        assert False, 'ART-2 не найден на листе Поставки - предложение'

    def test_burning_fill_when_lost_positive(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        """ART-1 имеет lost_orders=3.4 → ячейка «Упущено заказов» залита красным."""
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
        )
        wb = load_workbook(path)
        ws = wb['Поставки - предложение']
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == 'ART-1':
                lost_cell = ws.cell(row=row, column=6)
                assert lost_cell.value == 3
                assert lost_cell.fill.fgColor.rgb is not None
                assert 'FFCC' in (lost_cell.fill.fgColor.rgb or '')
                return
        assert False, 'ART-1 не найден на листе Поставки - предложение'

    def test_lost_below_half_shows_none_no_fill(self, warehouse_rows, tmp_path, monkeypatch):
        """lost_orders=0.3 → round=0 → None и без красной заливки."""
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        rows = [{
            'nm_id': 500, 'supplier_article': 'ART-LOW', 'barcode': '2000000000500',
            'subject': 'X', 'category': 'Y', 'product_group': 'A',
            'stock_qty': 5, 'in_way_from_client': 0, 'stock_qty_clean': 5,
            'orders_7d': 2, 'orders_14d': 4, 'orders_30d': None,
            'avg_per_day': 0.3, 'days_remaining': 16.0,
            'price_increase_pct': 10, 'price': 100.0,
            'availability': 'balanced', 'sale_rate_days': 16.0,
            'office_missing_days': 0.0, 'lost_orders': 0.3, 'trend_pct': 0.0,
        }]
        path = generate_report_from_data(rows, store_name='T', warehouse_rows=[])
        wb = load_workbook(path)
        ws = wb['Поставки - предложение']
        lost_cell = ws.cell(row=2, column=6)
        assert lost_cell.value is None
        rgb = lost_cell.fill.fgColor.rgb or ''
        assert 'FFCC' not in rgb

    def test_wb_suggestion_trend_fill(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        """ART-1 имеет trend_pct=15 → ячейка «Тренд» залита зелёным."""
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_report_from_data(
            product_rows, store_name='Test', warehouse_rows=warehouse_rows,
        )
        wb = load_workbook(path)
        ws = wb['Поставки - предложение']
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == 'ART-1':
                trend_cell = ws.cell(row=row, column=8)
                rgb = trend_cell.fill.fgColor.rgb or ''
                assert 'C8E6C9' in rgb
                return
        assert False, 'ART-1 не найден'


class TestComparisonReport:
    def test_generates_file(self, product_rows, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        store1 = product_rows[:2]
        store2 = [product_rows[0], product_rows[2]]
        path = generate_comparison_report(store1, store2, 'Магазин 1', 'Магазин 2')
        assert path is not None
        assert os.path.exists(path)

    def test_returns_none_if_no_common(self, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        store1 = [{'nm_id': 1, 'supplier_article': 'UNIQUE-1', 'stock_qty': 10,
                    'stock_qty_clean': 10, 'in_way_from_client': 0,
                    'avg_per_day': 1.0, 'price': 100}]
        store2 = [{'nm_id': 2, 'supplier_article': 'UNIQUE-2', 'stock_qty': 5,
                    'stock_qty_clean': 5, 'in_way_from_client': 0,
                    'avg_per_day': 2.0, 'price': 200}]
        path = generate_comparison_report(store1, store2, 'S1', 'S2')
        assert path is None

    def test_common_articles_counted(self, product_rows, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        path = generate_comparison_report(
            product_rows[:2], product_rows[:2], 'S1', 'S2',
        )
        wb = load_workbook(path)
        ws = wb.active
        data_rows = [r for r in ws.iter_rows(min_row=2) if r[1].value is not None]
        assert len(data_rows) == 4


class TestSummaryReport:
    def test_generates_file(self, product_rows, warehouse_rows, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        all_stores = {
            'Магазин 1': product_rows[:2],
            'Магазин 2': [product_rows[0], product_rows[2]],
        }
        all_wh = {
            'Магазин 1': warehouse_rows[:3],
            'Магазин 2': [warehouse_rows[0], warehouse_rows[3]],
        }
        path = generate_summary_report(all_stores, all_wh)
        assert path is not None
        assert os.path.exists(path)

    def test_returns_none_if_no_common(self, tmp_path, monkeypatch):
        monkeypatch.setattr('bot.config.REPORTS_DIR', str(tmp_path))
        store1 = [{'nm_id': 1, 'supplier_article': 'UNIQUE-1', 'stock_qty': 10,
                    'stock_qty_clean': 10, 'in_way_from_client': 0,
                    'avg_per_day': 1.0, 'price': 100}]
        store2 = [{'nm_id': 2, 'supplier_article': 'UNIQUE-2', 'stock_qty': 5,
                    'stock_qty_clean': 5, 'in_way_from_client': 0,
                    'avg_per_day': 2.0, 'price': 200}]
        path = generate_summary_report(
            {'S1': store1, 'S2': store2}, {'S1': [], 'S2': []},
        )
        assert path is None
